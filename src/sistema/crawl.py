from __future__ import annotations

from .browser_config import FirefoxConfig

from src.bootstrap import Directory
from src.db.tables import (
    BrowserContext as _BrowserContext,
    BrowserContextDict,
    ClientCertificate as _ClientCertificate,
    ColorSchemeType,
    CookieDict,
    ForcedColorsType,
    LocalStorageDict,
    LocaleType,
    OriginDict,
    ReducedMotionType,
    TimezoneIdType,
    Workbook as _Workbook,
    Worksheet as _Worksheet,
)
from src.exc import Task
from src.runtime import CommandLineArguments
from src.types import BrowserType

import gc
import asyncio
import hashlib
import inspect
import itertools

from asyncio import Lock, Queue, Semaphore
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Coroutine, Final, TypedDict
from urllib.parse import (
    ParseResult as URLParseResult,
    urlparse,
)

import playwright.async_api as playwright

from openpyxl import load_workbook


ESOCIAL_URL: Final[URLParseResult]  # TODO
GOVBR_URL: Final[URLParseResult]  # TODO


def make_db_origin_dict() -> OriginDict: ...


def make_db_cookie_dict() -> CookieDict: ...


def make_db_local_storage_dict() -> LocalStorageDict: ...


class CookieChangeDetail(TypedDict):
    oldCookie: str
    newCookie: str


def browserhook_on_cookie_change(detail: CookieChangeDetail) -> None: ...


class LocalStorageState(TypedDict):
    length: int
    keys: list[str]
    storageContent: dict[str, Any]


class LocalStorageDetail(TypedDict):
    oldLocalStorage: str  # strigfied LocalStorageState
    newLocalStorage: str  # strigfied LocalStorageState


def browserhook_on_local_storage_change(detail: LocalStorageDetail) -> None: ...


@dataclass(init=True, slots=True)
class PageState:
    origin_id: int | None
    cookie_ids: set[int]
    local_storage_ids: set[int]
    page_init_url: URLParseResult
    page: playwright.Page

    def state_hash(self) -> int:
        return (
            hash(self.origin_id) + hash(self.cookie_ids) + hash(self.local_storage_ids)
        )


class TaskInitState(TypedDict):
    sheet_path: Path
    certificate_db_id: int


class BrowserContext:
    # db properties
    accept_downloads: bool
    offline: bool
    javascript_enabled: bool
    is_mobile: bool
    has_touch: bool
    colorscheme: ColorSchemeType
    reduced_motion: ReducedMotionType
    forced_colors: ForcedColorsType
    locale: LocaleType
    timezone_id: TimezoneIdType

    def __init__(self, p: playwright.Playwright, browser: playwright.Browser) -> None:
        self.p = p
        self.browser = browser
        self.browser_exe = Path(browser.browser_type.executable_path)
        self.browser_type = BrowserType.enum_from_name(browser.browser_type.name)
        self._db_data_cache: dict[str, Any] = {}
        self._db_data_changes: dict[str, Any] = {}

    @staticmethod
    def default_db_dict() -> BrowserContextDict:
        return {
            'accept_downloads': False,
            'offline': False,
            'javascript_enabled': True,
            'is_mobile': False,
            'has_touch': False,
            'colorscheme': 'light',
            'reduced_motion': 'reduce',
            'forced_colors': 'none',
            'locale': 'pt-BR',
            'timezone_id': 'America/Sao_Paulo',
        }

    def _create_db_data_populate_self(self, init_state: TaskInitState):
        # caching init data
        self.sheet_path = init_state['sheet_path']
        self.certificate_db_id = init_state['certificate_db_id']

        # insert browser context data into db
        db_data = self.default_db_dict()
        self.browsercontext_id: int = _BrowserContext.sync_insert_one(db_data)

        # save data as in db
        inserted_db_data = _BrowserContext.sync_select_one_from_id(
            self.browsercontext_id
        )
        if inserted_db_data is None:
            raise RuntimeError('data not inserted')
        self._db_data_cache = inserted_db_data._asdict()

        # insert workbook data into db
        workbook_data = _Workbook.from_file(self.sheet_path)
        self.workbook_id: int = _Workbook.sync_insert_one(workbook_data)

        # insert worksheets data into db
        book = load_workbook(self.sheet_path, read_only=True)
        sheet_ids = []
        for _sheet in book.worksheets:
            data = _Worksheet.from_sheet_obj(book, _sheet)
            data['workbook_id'] = self.workbook_id
            _id = _Worksheet.sync_insert_one(data)
            sheet_ids.append(_id)
        self.worksheet_ids = tuple(sheet_ids)

        # explicity signal those should be collected when possible
        del db_data, inserted_db_data, workbook_data, book, sheet_ids, data, _id

    def _make_playwright_context_args(self) -> dict[str, Any]:
        indb = self._db_data_cache
        cert = _ClientCertificate.sync_select_one_from_id(self.certificate_db_id)
        cert_dict = None
        if cert is not None:
            cert_dict = {'origin': cert.origin}
            if cert.using_type == 'PFX':
                cert_dict['pfx'] = cert.pfx if cert.pfx is not None else b''
            elif cert.usign_type == 'KEY':
                cert_dict['key'] = cert.key if cert.key is not None else b''
            else:
                cert_dict['cert'] = cert.cert if cert.cert is not None else b''
            if cert.passphrase is not None:
                cert_dict['passphrase'] = cert.passphrase
        return dict(
            java_script_enabled=indb['javascript_enabled'],
            locale=indb['locale'],
            timezone_id=indb['timezone_id'],
            offline=indb['offline'],
            is_mobile=indb['is_mobile'],
            has_touch=indb['has_touch'],
            color_scheme=indb['colorscheme'],
            reduced_motion=indb['reduced_motion'],
            forced_colors=indb['forced_colors'],
            accept_downloads=indb['accept_downloads'],
            client_certificates=[cert_dict] if cert_dict is not None else None,
        )

    async def start_from(self, init_state: TaskInitState):
        self._create_db_data_populate_self(init_state)
        args = self._make_playwright_context_args()
        self.browser_context: playwright.BrowserContext = (
            await self.browser.new_context(**args)
        )

        # expose functions first
        await self.browser_context.expose_function(
            '_localStorageChangeHandler', browserhook_on_local_storage_change
        )
        await self.browser_context.expose_function(
            '_cookieChangeHandler', browserhook_on_cookie_change
        )

        # add scripts after exposing functions
        await self.browser_context.add_init_script(
            path=Path('./js/meta/cookieChange.js')
        )
        await self.browser_context.add_init_script(
            path=Path('./js/meta/localStorageChange.js')
        )

        # sync free memory
        gc.collect()

    def __getattr__(self, name: str, /) -> Any:
        if name in self._db_data_cache:
            return self._db_data_cache[name]
        return getattr(super(), name)

    def __setattr__(self, name: str, value: Any, /) -> None:
        if name in self._db_data_cache:
            self._db_data_changes[name] = value
            return
        setattr(super(), name, value)

    def can_update_db_data(self) -> bool:
        return len(self._db_data_changes) > 0

    def update_db_data(self):
        if (
            len(self._db_data_changes) == 0
            or len(self._db_data_cache) == 0
            or '_id' not in self._db_data_cache
        ):
            return
        _BrowserContext.sync_update_one_from_id(
            self._db_data_changes, self._db_data_cache['_id']
        )
        self._db_data_changes.clear()

    @staticmethod
    def firefox_exe() -> Path | None:
        path: Path | None = None
        try:
            gen = Directory.PLAYWRIGHT.glob('firefox-*/firefox/firefox.exe')
            path = next(gen)
            gen.close()
        except (GeneratorExit, StopIteration):
            pass
        return path

    async def new_page(self) -> PageState:
        parsed_url = urlparse()  # TODO
        page = await self.browser_context.new_page()
        return PageState(0, set(), set(), parsed_url, page)


StepFunc = Callable[['CrawlerTask'], None]


class CrawlerTask:
    _get_task_i = itertools.count(0).__next__
    tasks: dict[str, CrawlerTask] = {}
    _steps: tuple[StepFunc, ...]

    def __init__(
        self, p: playwright.Playwright, context: BrowserContext, runtime: BrowserRuntime
    ) -> None:
        self.p = p
        self.context = context
        self.runtime = runtime

    async def _crawler_run(self):
        await self.runtime.semaphore.acquire()
        self.page: PageState = await self.context.new_page()
        await self.runtime.task_count_increase()
        for step in self._steps:
            if self.context.can_update_db_data():
                self.context.update_db_data()
            if not getattr(step, '__crawler_step__', False):
                continue
            if inspect.iscoroutinefunction(step):
                await step(self)
                continue
            if inspect.isfunction(step):
                step(self)

    def schedule_self(self, *, return_if_cant: bool = False) -> None:
        if return_if_cant and not self.runtime.can_task_run_immediately():
            return
        self.task_index: int = self._get_task_i()
        self.task_index_formated: str = str(self.task_index).ljust(3, '0')
        self.task_name: str = f'crawler({self.task_index_formated})'
        self.asyncio_coro: Coroutine = self._crawler_run()
        self.asyncio_task: asyncio.Task = asyncio.create_task(
            self.asyncio_coro, name=self.task_name
        )
        self.asyncio_task.add_done_callback(self._asyncio_task_done_hook)

    def _asyncio_task_done_hook(self, task: asyncio.Task[None]) -> None:
        self.runtime.semaphore.release()
        self.unregister_self()

    def register_self(self):
        self.task_id: str = hashlib.md5(hash(self).to_bytes()).hexdigest().upper()
        self.tasks[self.task_id] = self

    def unregister_self(self):
        if hasattr(self, 'task_id') and self.task_id in self.tasks:
            del self.tasks[self.task_id]

    def __hash__(self) -> int:
        non_unique = hash(self.context.browser_exe) + hash(
            self.context.browser_type.playwright_name()
        )
        unique = hash(self.task_index)
        return non_unique + unique

    @classmethod
    def define_steps_in_strict_order(cls, *steps: StepFunc) -> None:
        if hasattr(cls, '_steps'):
            return
        if any(not hasattr(func, '__crawler_step__') for func in steps):
            raise Task.CannotRegisterStep(
                'function to register was not defined as a CrawlerTask.step'
            )
        cls._steps = tuple(steps)

    @classmethod
    def step(cls, func: StepFunc) -> StepFunc:
        func.__crawler_step__ = True
        return func


class BrowserRuntime:
    cli_args = CommandLineArguments().parse_argv()
    SEMAPHORE_LIMIT: Final[int] = 5
    browser: playwright.Browser

    def __init__(self, p: playwright.Playwright) -> None:
        self.p = p
        self.semaphore: Semaphore = Semaphore(self.SEMAPHORE_LIMIT)
        self.sheet_queue: Queue[TaskInitState] = Queue()
        self._lifetime_task_count: int = 0
        self._task_count_lock = Lock()

    @property
    def lifetime_task_count(self) -> int:
        return self._lifetime_task_count

    async def task_count_increase(self):
        async with self._task_count_lock:
            self._lifetime_task_count += 1

    @classmethod
    def should_run(cls) -> bool:
        return not cls.cli_args.no_playwright

    def can_task_run_immediately(self) -> bool:
        return not (self.sheet_queue.empty() and self.semaphore.locked())

    def can_schedule_task(self) -> bool:
        return not self.sheet_queue.empty()

    async def new_firefox(self) -> playwright.Browser:
        exe = BrowserContext.firefox_exe()
        if exe is None:
            raise RuntimeError('cannot find firefox executable')

        readwrite = 0o666
        readonly = 0o444

        # firefox distribution config dir
        dist_dir = exe.parent / 'distribution'
        if not dist_dir.exists():
            dist_dir.mkdir(parents=True, exist_ok=True)

        # check/write override.ini
        override_path = dist_dir / 'override.ini'
        if not override_path.exists():
            override_path.touch(readwrite, exist_ok=True)
        file_hash = FirefoxConfig.hash_override_ini(ini_file=override_path)
        dict_hash = FirefoxConfig.hash_override_ini()
        if file_hash != dict_hash:
            if override_path.stat().st_mode != readwrite:
                override_path.chmod(readwrite)
            override_path.write_bytes(FirefoxConfig.parse_override_ini())
        if override_path.stat().st_mode != readonly:
            override_path.chmod(readonly)

        # check/write policies.json
        policies_path = dist_dir / 'policies.json'
        if policies_path.exists():
            policies_path.touch(readwrite, exist_ok=True)
        file_hash = FirefoxConfig.hash_policies_json(policies_file=policies_path)
        dict_hash = FirefoxConfig.hash_policies_json()
        if file_hash != dict_hash:
            if policies_path.stat().st_mode != readwrite:
                policies_path.chmod(readwrite)
            policies_path.write_bytes(FirefoxConfig.parse_policies_json())
        if policies_path.stat().st_mode != readonly:
            policies_path.chmod(readonly)

        return await self.p.firefox.launch(
            executable_path=exe,
            args=['-override', f'"{str(override_path)}"'],
            headless=not __debug__,
            # devtools disabled in policies.json
            chromium_sandbox=False,
            # downloads_path and traces_dir are currently not, but can be set in policies.json
            downloads_path=Directory.BROWSER_DOWNLOADS,
            traces_dir=Directory.BROWSER_TRACES,
            firefox_user_prefs=FirefoxConfig.user_js,
        )

    async def schedule_task(self):
        if self.sheet_queue.empty():
            return
        if not hasattr(self, 'browser'):
            self.browser = await self.new_firefox()
        context = BrowserContext(self.p, self.browser)
        init_state = await self.sheet_queue.get()
        await context.start_from(init_state)
        task = CrawlerTask(self.p, context, self)
        task.register_self()
        task.schedule_self()


CrawlerTask.define_steps_in_strict_order  # TODO
